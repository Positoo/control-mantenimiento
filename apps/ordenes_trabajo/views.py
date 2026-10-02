from django.db.models import Q
from django.shortcuts import get_object_or_404, render, redirect
from django.http import JsonResponse
from django.contrib.auth.decorators import permission_required
from django.db import transaction

from .models import DetalleOrdenTrabajo, OrdenTrabajo
from .forms import DetalleOrdenTrabajoFormSet, OrdenTrabajoForm


from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse

def inicio(request):
    ordenes = OrdenTrabajo.objects.all()

    busqueda = request.GET.get("buscar", "").strip()
    estado = request.GET.get("estado", "").strip()
    tipo = request.GET.get("tipo", "").strip()
    fecha_desde = request.GET.get("fecha_desde", "").strip()
    fecha_hasta = request.GET.get("fecha_hasta", "").strip()

    if busqueda:
        ordenes = ordenes.filter(
            Q(movil__icontains=busqueda)
            | Q(dominio__icontains=busqueda)
            | Q(id__icontains=busqueda)
        )

    if estado == "abierta":
        ordenes = ordenes.filter(fecha_egreso__isnull=True)

    elif estado == "cerrada":
        ordenes = ordenes.filter(fecha_egreso__isnull=False)

    if tipo == "preventivo":
        ordenes = ordenes.filter(
            detalles__tipo=DetalleOrdenTrabajo.Tipo.PREVENTIVO
        ).distinct()

    elif tipo == "correctivo":
        ordenes = ordenes.filter(
            detalles__tipo=DetalleOrdenTrabajo.Tipo.CORRECTIVO
        ).distinct()

    if fecha_desde:
        ordenes = ordenes.filter(fecha_ingreso__gte=fecha_desde)

    if fecha_hasta:
        ordenes = ordenes.filter(fecha_ingreso__lte=fecha_hasta)

    return render(
        request,
        "ordenes_trabajo/lista.html",
        {
            "ordenes": ordenes,
            "busqueda": busqueda,
            "estado": estado,
            "tipo": tipo,
            "fecha_desde": fecha_desde,
            "fecha_hasta": fecha_hasta,
        },
    )


def reporte_costos(request):
    from collections import defaultdict
    from apps.unidades.models import UnidadDeNegocio, Alcance

    ordenes = OrdenTrabajo.objects.all()

    fecha_desde = request.GET.get("fecha_desde", "").strip()
    fecha_hasta = request.GET.get("fecha_hasta", "").strip()
    unidad_negocio = request.GET.get("unidad_negocio", "").strip()
    alcance = request.GET.get("alcance", "").strip()
    movil = request.GET.get("movil", "").strip()
    tipo = request.GET.get("tipo", "").strip()

    # Filtros
    if fecha_desde:
        ordenes = ordenes.filter(fecha_ingreso__gte=fecha_desde)

    if fecha_hasta:
        ordenes = ordenes.filter(fecha_ingreso__lte=fecha_hasta)

    if unidad_negocio:
        ordenes = ordenes.filter(
            unidad_de_negocio_id=unidad_negocio
        )

    if alcance:
        ordenes = ordenes.filter(
            alcance_id=alcance
        )

    if movil:
        ordenes = ordenes.filter(
            movil__icontains=movil
        )

    if tipo == "preventivo":
        ordenes = ordenes.filter(
            detalles__tipo=DetalleOrdenTrabajo.Tipo.PREVENTIVO
        ).distinct()

    elif tipo == "correctivo":
        ordenes = ordenes.filter(
            detalles__tipo=DetalleOrdenTrabajo.Tipo.CORRECTIVO
        ).distinct()

    # Estructura:
    # Unidad de Negocio → Alcance → Unidad
    reporte = defaultdict(
        lambda: defaultdict(
            lambda: defaultdict(
                lambda: {
                    "preventivo": 0,
                    "correctivo": 0,
                    "total": 0,
                }
            )
        )
    )

    for ot in ordenes:
        detalles = ot.detalles.all()

        for detalle in detalles:
            costo = detalle.costo_total

            if detalle.tipo == DetalleOrdenTrabajo.Tipo.PREVENTIVO:
                reporte[
                    ot.unidad_de_negocio
                ][
                    ot.alcance
                ][
                    ot.unidad
                ]["preventivo"] += costo

            elif detalle.tipo == DetalleOrdenTrabajo.Tipo.CORRECTIVO:
                reporte[
                    ot.unidad_de_negocio
                ][
                    ot.alcance
                ][
                    ot.unidad
                ]["correctivo"] += costo

            reporte[
                ot.unidad_de_negocio
            ][
                ot.alcance
            ][
                ot.unidad
            ]["total"] += costo

    # Convertimos defaultdict en una estructura más cómoda para el template.
    reporte_final = []

    total_general_preventivo = 0
    total_general_correctivo = 0
    total_general = 0

    for negocio, alcances in reporte.items():

        negocio_preventivo = 0
        negocio_correctivo = 0
        negocio_total = 0

        alcances_final = []

        for alcance_obj, unidades in alcances.items():

            alcance_preventivo = 0
            alcance_correctivo = 0
            alcance_total = 0

            unidades_final = []

            for unidad, costos in unidades.items():

                unidad_total = costos["total"]

                unidades_final.append({
                    "unidad": unidad,
                    "preventivo": costos["preventivo"],
                    "correctivo": costos["correctivo"],
                    "total": unidad_total,
                })

                alcance_preventivo += costos["preventivo"]
                alcance_correctivo += costos["correctivo"]
                alcance_total += unidad_total

            alcances_final.append({
                "alcance": alcance_obj,
                "preventivo": alcance_preventivo,
                "correctivo": alcance_correctivo,
                "total": alcance_total,
                "unidades": unidades_final,
            })

            negocio_preventivo += alcance_preventivo
            negocio_correctivo += alcance_correctivo
            negocio_total += alcance_total

        reporte_final.append({
            "negocio": negocio,
            "preventivo": negocio_preventivo,
            "correctivo": negocio_correctivo,
            "total": negocio_total,
            "alcances": alcances_final,
        })

        total_general_preventivo += negocio_preventivo
        total_general_correctivo += negocio_correctivo
        total_general += negocio_total

    unidades_de_negocio = UnidadDeNegocio.objects.filter(
        activo=True
    ).order_by("nombre")

    alcances = Alcance.objects.filter(
        activo=True
    ).order_by("nombre")

    return render(
        request,
        "ordenes_trabajo/reporte_costos.html",
        {
            "reporte": reporte_final,
            "total_general_preventivo": total_general_preventivo,
            "total_general_correctivo": total_general_correctivo,
            "total_general": total_general,
            "unidades_de_negocio": unidades_de_negocio,
            "alcances": alcances,
            "fecha_desde": fecha_desde,
            "fecha_hasta": fecha_hasta,
            "unidad_negocio": unidad_negocio,
            "alcance": alcance,
            "movil": movil,
            "tipo": tipo,
        },
    )


def detalle_ot(request, pk):
    ot = get_object_or_404(
        OrdenTrabajo,
        pk=pk,
    )

    return render(
        request,
        "ordenes_trabajo/detalle.html",
        {
            "ot": ot,
        },
    )

@permission_required(
    "ordenes_trabajo.add_ordentrabajo",
    raise_exception=True,
)
def nueva_ot(request):
    if request.method == "POST":
        form = OrdenTrabajoForm(request.POST)
        formset = DetalleOrdenTrabajoFormSet(request.POST)

        if form.is_valid() and formset.is_valid():
            ot = form.save(commit=False)

            ot.unidad = form.cleaned_data["unidad_busqueda"]
            ot.movil = ot.unidad.movil or ""
            ot.dominio = ot.unidad.dominio
            ot.unidad_de_negocio = ot.unidad.unidad_de_negocio
            ot.alcance = ot.unidad.alcance

            ot.save()

            detalles = formset.save(commit=False)

            for detalle in detalles:
                detalle.orden_trabajo = ot
                detalle.save()

            return redirect(
                "detalle_ot",
                pk=ot.pk,
            )

    else:
        form = OrdenTrabajoForm()
        formset = DetalleOrdenTrabajoFormSet()

    return render(
        request,
        "ordenes_trabajo/nueva.html",
        {
            "form": form,
            "formset": formset,
        },
    )


@permission_required(
    "ordenes_trabajo.change_ordentrabajo",
    raise_exception=True,
)
def editar_ot(request, pk):
    ot = get_object_or_404(
        OrdenTrabajo,
        pk=pk,
    )

    if request.method == "POST":
        form = OrdenTrabajoForm(request.POST, instance=ot)
        formset = DetalleOrdenTrabajoFormSet(
            request.POST,
            instance=ot,
        )

        if form.is_valid() and formset.is_valid():
            with transaction.atomic():

                ot = form.save(commit=False)

                unidad = form.cleaned_data["unidad_busqueda"]

                ot.unidad = unidad
                ot.movil = unidad.movil or ""
                ot.dominio = unidad.dominio
                ot.unidad_de_negocio = unidad.unidad_de_negocio
                ot.alcance = unidad.alcance

                ot.save()

                formset.save()

            return redirect(
                "detalle_ot",
                pk=ot.pk,
            )

    else:
        form = OrdenTrabajoForm(
            instance=ot,
            initial={
                "unidad_busqueda": ot.dominio,
            },
        )

        formset = DetalleOrdenTrabajoFormSet(
            instance=ot,
        )

    return render(
        request,
        "ordenes_trabajo/editar.html",
        {
            "form": form,
            "formset": formset,
            "ot": ot,
        },
    )


def buscar_unidad(request):
    valor = request.GET.get("valor", "").strip()

    if not valor:
        return JsonResponse({
            "encontrada": False,
            "mensaje": "Ingrese un dominio o móvil.",
        })

    from apps.unidades.models import Unidad

    unidades = Unidad.objects.filter(
        activo=True
    ).filter(
        Q(dominio__iexact=valor) | Q(movil__iexact=valor)
    )

    if not unidades.exists():
        return JsonResponse({
            "encontrada": False,
            "mensaje": "No existe una unidad activa con ese dominio o móvil.",
        })

    if unidades.count() > 1:
        return JsonResponse({
            "encontrada": False,
            "mensaje": "Hay más de una unidad que coincide con ese valor.",
        })

    unidad = unidades.first()

    return JsonResponse({
        "encontrada": True,
        "id": unidad.id_interno,
        "dominio": unidad.dominio,
        "movil": unidad.movil or "",
    })

############## PARA EXPORTAR COMO EXCEL ############

def exportar_costos_excel(request):
    from collections import defaultdict
    from apps.unidades.models import UnidadDeNegocio, Alcance

    ordenes = OrdenTrabajo.objects.all()

    fecha_desde = request.GET.get("fecha_desde", "").strip()
    fecha_hasta = request.GET.get("fecha_hasta", "").strip()
    unidad_negocio = request.GET.get("unidad_negocio", "").strip()
    alcance = request.GET.get("alcance", "").strip()
    movil = request.GET.get("movil", "").strip()
    tipo = request.GET.get("tipo", "").strip()

    # Filtros
    if fecha_desde:
        ordenes = ordenes.filter(fecha_ingreso__gte=fecha_desde)

    if fecha_hasta:
        ordenes = ordenes.filter(fecha_ingreso__lte=fecha_hasta)

    if unidad_negocio:
        ordenes = ordenes.filter(
            unidad_de_negocio_id=unidad_negocio
        )

    if alcance:
        ordenes = ordenes.filter(
            alcance_id=alcance
        )

    if movil:
        ordenes = ordenes.filter(
            movil__icontains=movil
        )

    if tipo == "preventivo":
        ordenes = ordenes.filter(
            detalles__tipo=DetalleOrdenTrabajo.Tipo.PREVENTIVO
        ).distinct()

    elif tipo == "correctivo":
        ordenes = ordenes.filter(
            detalles__tipo=DetalleOrdenTrabajo.Tipo.CORRECTIVO
        ).distinct()

    # Estructura:
    # Unidad de Negocio → Alcance → Unidad
    reporte = defaultdict(
        lambda: defaultdict(
            lambda: defaultdict(
                lambda: {
                    "preventivo": 0,
                    "correctivo": 0,
                    "total": 0,
                }
            )
        )
    )

    for ot in ordenes:
        for detalle in ot.detalles.all():

            costo = detalle.costo_total

            if detalle.tipo == DetalleOrdenTrabajo.Tipo.PREVENTIVO:
                reporte[
                    ot.unidad_de_negocio
                ][
                    ot.alcance
                ][
                    ot.unidad
                ]["preventivo"] += costo

            elif detalle.tipo == DetalleOrdenTrabajo.Tipo.CORRECTIVO:
                reporte[
                    ot.unidad_de_negocio
                ][
                    ot.alcance
                ][
                    ot.unidad
                ]["correctivo"] += costo

            reporte[
                ot.unidad_de_negocio
            ][
                ot.alcance
            ][
                ot.unidad
            ]["total"] += costo

    # Crear Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Costos de mantenimiento"

    # Título
    ws["A1"] = "Costos de mantenimiento"
    ws["A1"].font = Font(bold=True, size=16)

    ws.merge_cells("A1:F1")

    # Filtros aplicados
    fila = 3

    ws.cell(fila, 1, "Filtros aplicados")
    ws.cell(fila, 1).font = Font(bold=True)

    fila += 1

    filtros = [
        ("Fecha desde", fecha_desde or "Todas"),
        ("Fecha hasta", fecha_hasta or "Todas"),
        ("Unidad de Negocio", unidad_negocio or "Todas"),
        ("Alcance", alcance or "Todos"),
        ("Móvil", movil or "Todos"),
        ("Tipo", tipo or "Todos"),
    ]

    for nombre, valor in filtros:
        ws.cell(fila, 1, nombre)
        ws.cell(fila, 2, valor)
        fila += 1

    fila += 1

    # Encabezados
    encabezados = [
        "Unidad de Negocio",
        "Alcance",
        "Móvil",
        "Dominio",
        "Preventivo",
        "Correctivo",
        "Total",
    ]

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(fila, columna, encabezado)
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")

    fila += 1

    # Datos
    for negocio, alcances in reporte.items():

        for alcance_obj, unidades in alcances.items():

            for unidad, costos in unidades.items():

                ws.cell(fila, 1, negocio.nombre)
                ws.cell(fila, 2, alcance_obj.nombre)
                ws.cell(fila, 3, unidad.movil or "")
                ws.cell(fila, 4, unidad.dominio)

                ws.cell(fila, 5, float(costos["preventivo"]))
                ws.cell(fila, 6, float(costos["correctivo"]))
                ws.cell(fila, 7, float(costos["total"]))

                fila += 1

    # Formato monetario
    for fila_excel in range(1, fila):
        for columna in (5, 6, 7):
            ws.cell(fila_excel, columna).number_format = '$ #,##0.00'

    # Ancho de columnas
    anchos = {
        "A": 25,
        "B": 25,
        "C": 15,
        "D": 15,
        "E": 18,
        "F": 18,
        "G": 18,
    }

    for columna, ancho in anchos.items():
        ws.column_dimensions[columna].width = ancho

    # Respuesta HTTP
    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="costos_mantenimiento.xlsx"'
    )

    wb.save(response)

    return response

############## FIN EXPORTACION #####################